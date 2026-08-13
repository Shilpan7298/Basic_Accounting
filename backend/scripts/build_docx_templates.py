"""Build the starter Word templates.

These are *starters*. The intended end state is that Urjapod's own existing
invoice and PO .docx files become the templates — you open one in Word, drop
the ``{{ }}`` tags in, save it as ``<name>-v<N>.docx``, and the app picks it up
with no code change. This script exists so there is something to render on day
one, and so the tag syntax has a worked example to copy.

Run:  python backend/scripts/build_docx_templates.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1] / "app" / "templates" / "documents"


def _set_cell(cell, text, *, bold=False, size=8.5, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return cell


def _para(doc, text, *, bold=False, size=10, align=None, color=None, space_after=4):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return paragraph


def _row_loop(table, cell_texts, *, right_align=frozenset(), var="line", over="lines"):
    """Wire up a docxtpl table-row loop across the table's last three rows.

    docxtpl's ``{%tr %}`` tag *consumes the row it sits in* — the regex that
    lifts the tag out replaces the whole ``<w:tr>`` with the bare Jinja tag. So
    ``for`` and ``endfor`` each need a row to themselves, with the repeating
    content row sandwiched between them. Putting both tags in the one content
    row deletes that row and leaves a dangling ``endfor``.
    """
    open_row, body_row, close_row = table.rows[-3], table.rows[-2], table.rows[-1]
    _set_cell(open_row.cells[0], "{%%tr for %s in %s %%}" % (var, over), size=1)
    for index, text in enumerate(cell_texts):
        _set_cell(
            body_row.cells[index],
            text,
            size=8,
            align=WD_ALIGN_PARAGRAPH.RIGHT if index in right_align else None,
        )
    _set_cell(close_row.cells[0], "{%tr endfor %}", size=1)


def _narrow_margins(doc):
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, attr, Pt(36))


def build_invoice(path: Path) -> None:
    doc = Document()
    _narrow_margins(doc)

    _para(doc, "{{ company.legal_name }}", bold=True, size=15)
    _para(doc, "{% for l in company.address_lines %}{{ l }}{% if not loop.last %}, {% endif %}{% endfor %}", size=8, color=(0x55, 0x55, 0x55))
    _para(doc, "GSTIN: {{ company.gstin }}  |  PAN: {{ company.pan }}  |  {{ company.phone }}  |  {{ company.email }}",
          size=8, color=(0x55, 0x55, 0x55), space_after=10)

    _para(doc, "{{ document.title }}", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "{{ document.subtitle }}", size=8, align=WD_ALIGN_PARAGRAPH.CENTER,
          color=(0x44, 0x44, 0x44), space_after=8)

    # The proforma disclaimer only prints on a proforma.
    _para(doc,
          "{% if document.is_proforma %}This is NOT a tax invoice. No GST liability arises on this "
          "document. It is issued to request payment against the milestone shown below."
          "{% endif %}",
          size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x6B, 0x4B, 0x00), space_after=8)

    meta = doc.add_table(rows=2, cols=4)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_cell(meta.cell(0, 0), "Number", bold=True)
    _set_cell(meta.cell(0, 1), "{{ document.number }}")
    _set_cell(meta.cell(0, 2), "Date", bold=True)
    _set_cell(meta.cell(0, 3), "{{ document.date }}")
    _set_cell(meta.cell(1, 0), "Customer PO", bold=True)
    _set_cell(meta.cell(1, 1), "{{ document.po_reference }}")
    _set_cell(meta.cell(1, 2), "Place of Supply", bold=True)
    _set_cell(meta.cell(1, 3), "{{ party.state_name }} ({{ party.state_code }})")

    doc.add_paragraph()
    parties = doc.add_table(rows=1, cols=2)
    parties.style = "Table Grid"
    _set_cell(
        parties.cell(0, 0),
        "BILL TO\n{{ party.name }}\n"
        "{% for l in party.billing_address_lines %}{{ l }}\n{% endfor %}"
        "GSTIN: {{ party.gstin }}",
    )
    _set_cell(
        parties.cell(0, 1),
        "SHIP TO\n{{ party.name }}\n"
        "{% for l in party.shipping_address_lines %}{{ l }}\n{% endfor %}",
    )

    doc.add_paragraph()
    lines = doc.add_table(rows=4, cols=9)
    lines.style = "Table Grid"
    headers = ["#", "Description", "HSN", "Qty", "UOM", "Rate", "Taxable", "GST%", "Amount"]
    for index, header in enumerate(headers):
        _set_cell(lines.cell(0, index), header, bold=True, size=8,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    body = [
        "{{ line.no }}",
        "{{ line.description }}",
        "{{ line.hsn_code }}",
        "{{ line.qty.fmt }}",
        "{{ line.uom }}",
        "{{ line.unit_price.fmt }}",
        "{{ line.taxable_value.fmt }}",
        "{{ line.gst_rate_percent.fmt }}",
        "{{ line.total.fmt }}",
    ]
    _row_loop(lines, body, right_align={3, 5, 6, 7, 8})

    doc.add_paragraph()
    totals = doc.add_table(rows=6, cols=2)
    totals.style = "Table Grid"
    rows = [
        ("Taxable Value", "{{ totals.taxable_value.fmt }}"),
        ("{% if tax.is_intra_state %}CGST{% else %}IGST{% endif %}",
         "{% if tax.is_intra_state %}{{ tax.cgst_total.fmt }}{% else %}{{ tax.igst_total.fmt }}{% endif %}"),
        ("{% if tax.is_intra_state %}SGST{% else %}Cess{% endif %}",
         "{% if tax.is_intra_state %}{{ tax.sgst_total.fmt }}{% else %}{{ tax.cess_total.fmt }}{% endif %}"),
        ("Round Off", "{{ totals.round_off.fmt }}"),
        ("GRAND TOTAL", "{{ totals.grand_total.fmt }}"),
        ("Balance Due", "{{ payment.balance_due.fmt }}"),
    ]
    for index, (label, value) in enumerate(rows):
        _set_cell(totals.cell(index, 0), label, bold=index >= 4)
        _set_cell(totals.cell(index, 1), value, bold=index >= 4,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)

    _para(doc, "Amount in words: {{ totals.amount_in_words }}", size=9, space_after=8)
    _para(doc,
          "{% if document.is_proforma %}Milestone: {{ payment.milestone_label }} "
          "({{ payment.milestone_percent.fmt }}% of order value){% endif %}", size=9)
    _para(doc, "Payment terms: {{ payment.terms_text }}", size=8.5, color=(0x55, 0x55, 0x55))
    _para(doc,
          "Bank: {{ company.bank_name }}, {{ company.bank_branch }} | "
          "A/c {{ company.bank_account_no }} | IFSC {{ company.bank_ifsc }}",
          size=8.5, space_after=16)

    _para(doc, "For {{ company.legal_name }}", size=9, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, "\n\nAuthorised Signatory", size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)
    _para(doc,
          "Generated {{ meta.generated_at_ist }} IST | template {{ meta.template_name }} "
          "v{{ meta.template_version }}", size=7, color=(0x88, 0x88, 0x88))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def build_purchase_order(path: Path) -> None:
    doc = Document()
    _narrow_margins(doc)

    _para(doc, "{{ company.legal_name }}", bold=True, size=15)
    _para(doc, "{% for l in company.address_lines %}{{ l }}{% if not loop.last %}, {% endif %}{% endfor %}",
          size=8, color=(0x55, 0x55, 0x55))
    _para(doc, "GSTIN: {{ company.gstin }}  |  {{ company.phone }}  |  {{ company.email }}",
          size=8, color=(0x55, 0x55, 0x55), space_after=10)

    _para(doc, "PURCHASE ORDER", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    meta = doc.add_table(rows=2, cols=4)
    meta.style = "Table Grid"
    _set_cell(meta.cell(0, 0), "PO Number", bold=True)
    _set_cell(meta.cell(0, 1), "{{ document.number }}")
    _set_cell(meta.cell(0, 2), "PO Date", bold=True)
    _set_cell(meta.cell(0, 3), "{{ document.date }}")
    _set_cell(meta.cell(1, 0), "Required By", bold=True)
    _set_cell(meta.cell(1, 1), "{{ document.due_date }}")
    _set_cell(meta.cell(1, 2), "Place of Supply", bold=True)
    _set_cell(meta.cell(1, 3), "{{ party.state_name }} ({{ party.state_code }})")

    doc.add_paragraph()
    parties = doc.add_table(rows=1, cols=2)
    parties.style = "Table Grid"
    _set_cell(
        parties.cell(0, 0),
        "SUPPLIER\n{{ party.name }}\n"
        "{% for l in party.billing_address_lines %}{{ l }}\n{% endfor %}"
        "GSTIN: {{ party.gstin }}",
    )
    _set_cell(
        parties.cell(0, 1),
        "DELIVER TO\n{{ document.delivery_location }}\n{{ company.legal_name }}\n"
        "{% for l in company.address_lines %}{{ l }}\n{% endfor %}",
    )

    doc.add_paragraph()
    lines = doc.add_table(rows=4, cols=8)
    lines.style = "Table Grid"
    for index, header in enumerate(
        ["#", "Description", "HSN", "Qty", "UOM", "Rate", "Taxable", "Amount"]
    ):
        _set_cell(lines.cell(0, index), header, bold=True, size=8,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    body = [
        "{{ line.no }}",
        "{{ line.description }}",
        "{{ line.hsn_code }}",
        "{{ line.qty.fmt }}",
        "{{ line.uom }}",
        "{{ line.unit_price.fmt }}",
        "{{ line.taxable_value.fmt }}",
        "{{ line.total.fmt }}",
    ]
    _row_loop(lines, body, right_align={3, 5, 6, 7})

    doc.add_paragraph()
    totals = doc.add_table(rows=5, cols=2)
    totals.style = "Table Grid"
    rows = [
        ("Taxable Value", "{{ totals.taxable_value.fmt }}"),
        ("{% if tax.is_intra_state %}CGST{% else %}IGST{% endif %}",
         "{% if tax.is_intra_state %}{{ tax.cgst_total.fmt }}{% else %}{{ tax.igst_total.fmt }}{% endif %}"),
        ("{% if tax.is_intra_state %}SGST{% else %}Cess{% endif %}",
         "{% if tax.is_intra_state %}{{ tax.sgst_total.fmt }}{% else %}{{ tax.cess_total.fmt }}{% endif %}"),
        ("Round Off", "{{ totals.round_off.fmt }}"),
        ("TOTAL", "{{ totals.grand_total.fmt }}"),
    ]
    for index, (label, value) in enumerate(rows):
        _set_cell(totals.cell(index, 0), label, bold=index == 4)
        _set_cell(totals.cell(index, 1), value, bold=index == 4,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)

    _para(doc, "Amount in words: {{ totals.amount_in_words }}", size=9, space_after=8)
    _para(doc, "Payment Terms: {{ terms.payment_terms_text }}", size=9)
    _para(doc, "Freight / Delivery: {{ terms.freight_terms }}", size=9)
    _para(doc, "Warranty: {{ terms.warranty_text }}", size=9)
    _para(doc, "Notes: {{ document.notes }}", size=9, space_after=10)

    _para(doc,
          "1. Please quote this PO number on all invoices, packing lists and correspondence.\n"
          "2. Goods not conforming to the specification above may be rejected at your cost.\n"
          "3. Invoices must carry the correct HSN code and GSTIN as printed on this order.",
          size=8, color=(0x55, 0x55, 0x55), space_after=16)

    _para(doc, "For {{ company.legal_name }}", size=9, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, "\n\nAuthorised Signatory", size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)
    _para(doc,
          "Generated {{ meta.generated_at_ist }} IST | template {{ meta.template_name }} "
          "v{{ meta.template_version }}", size=7, color=(0x88, 0x88, 0x88))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main() -> int:
    invoice_path = ROOT / "invoice" / "standard-v1.docx"
    po_path = ROOT / "purchase_order" / "standard-v1.docx"
    build_invoice(invoice_path)
    build_purchase_order(po_path)
    print(f"wrote {invoice_path}")
    print(f"wrote {po_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
