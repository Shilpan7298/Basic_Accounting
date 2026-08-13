"""The audit log and the template contract.

Both are constraints rather than features (CLAUDE.md #2 and #3), so the tests
are about what *cannot* happen: a silent mutation, and a template that
requires a code change.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from sqlalchemy import select

from app.models import AuditAction, AuditEvent, Customer
from app.rendering import render, templates
from app.rendering.context import GOLDEN_SAMPLE
from app.services import invoices as invoice_service
from app.services import orders as order_service


# --- audit -----------------------------------------------------------------

def _events(db, entity_type=None, entity_id=None):
    stmt = select(AuditEvent)
    if entity_type:
        stmt = stmt.where(AuditEvent.entity_type == entity_type)
    if entity_id:
        stmt = stmt.where(AuditEvent.entity_id == entity_id)
    return list(db.execute(stmt).scalars())


def test_creating_an_audited_row_is_logged_without_anyone_asking(db):
    """The catch-all hook: a mutation nobody logged explicitly is still logged."""
    customer = Customer(name="Silent Insert Ltd", gstin=None)
    db.add(customer)
    db.commit()

    events = _events(db, "customers", customer.id)
    assert events, "the session hook did not record the insert"
    assert events[0].action == AuditAction.CREATE
    assert events[0].after_json["name"] == "Silent Insert Ltd"


def test_updating_a_row_records_before_and_after(db, gujarat_customer):
    gujarat_customer.name = "Greenfield Textiles Pvt Ltd"
    db.commit()

    updates = [e for e in _events(db, "customers", gujarat_customer.id)
               if e.action in (AuditAction.UPDATE, AuditAction.STATUS_CHANGE)]
    assert updates
    latest = updates[-1]
    assert latest.before_json["name"] == "Greenfield Textiles Limited"
    assert latest.after_json["name"] == "Greenfield Textiles Pvt Ltd"


def test_only_changed_columns_appear_in_the_diff(db, gujarat_customer):
    gujarat_customer.phone = "+91 79 1111 2222"
    db.commit()

    latest = [e for e in _events(db, "customers", gujarat_customer.id)
              if e.action == AuditAction.UPDATE][-1]
    assert set(latest.after_json) == {"phone"}


def test_a_status_change_is_recorded_as_such(db, gujarat_customer, order_factory):
    order = order_factory(gujarat_customer)
    order_service.transition(db, order, "in_production", actor="alice")
    db.commit()

    changes = [e for e in _events(db, "sales_orders", order.id)
               if e.action == AuditAction.STATUS_CHANGE]
    assert changes
    assert any(
        e.before_json.get("status") == "received"
        and e.after_json.get("status") == "in_production"
        for e in changes
    )


def test_the_actor_is_recorded(db, gujarat_customer, order_factory):
    order = order_factory(gujarat_customer)
    order_service.transition(db, order, "in_production", actor="alice@urjapod")
    db.commit()

    actors = {e.actor for e in _events(db, "sales_orders", order.id)}
    assert "alice@urjapod" in actors


def test_issuing_an_invoice_writes_an_issue_event(
    db, company, tax_rates, gujarat_customer, order_factory
):
    order = order_factory(gujarat_customer)
    invoice = invoice_service.create_tax_invoice(db, order, actor="bob")
    invoice_service.issue(db, invoice, actor="bob")
    db.commit()

    actions = {e.action for e in _events(db, "invoices", invoice.id)}
    assert AuditAction.CREATE in actions
    assert AuditAction.ISSUE in actions


def test_cancelling_writes_a_cancel_event_with_the_reason(
    db, company, tax_rates, gujarat_customer, order_factory
):
    order = order_factory(gujarat_customer)
    invoice = invoice_service.create_tax_invoice(db, order, actor="bob")
    invoice_service.cancel(db, invoice, actor="bob", reason="duplicate")
    db.commit()

    cancels = [e for e in _events(db, "invoices", invoice.id) if e.action == AuditAction.CANCEL]
    assert cancels
    assert cancels[-1].context_json["reason"] == "duplicate"


def test_the_audit_log_is_append_only_in_the_api_surface():
    """There is no route that deletes audit rows. Assert it structurally."""
    from app.main import app

    audit_routes = [
        (route.path, method)
        for route in app.routes
        for method in getattr(route, "methods", []) or []
        if "audit" in getattr(route, "path", "")
    ]
    for path in app.openapi()["paths"]:
        if "audit" in path:
            assert set(app.openapi()["paths"][path]) <= {"get"}, path
    assert all(method in {"GET", "HEAD", "OPTIONS"} for _p, method in audit_routes)


def test_money_survives_the_audit_json_round_trip(
    db, company, tax_rates, gujarat_customer, order_factory
):
    """Decimals must not become floats on the way into the log."""
    order = order_factory(gujarat_customer, lines=(("x", "8507", "3", "3333.33"),))
    invoice = invoice_service.create_tax_invoice(db, order, actor="t")
    db.commit()

    created = [e for e in _events(db, "invoices", invoice.id) if e.action == AuditAction.CREATE]
    stored = Decimal(created[-1].after_json["grand_total"])
    assert stored == Decimal(invoice.grand_total)


# --- templates -------------------------------------------------------------

def test_templates_are_discovered_by_scanning_not_by_a_registry():
    found = templates.discover()
    labels = {(t.doc_type, t.name, t.version, t.kind) for t in found}

    assert ("invoice", "standard", 1, "html") in labels
    assert ("invoice", "standard", 1, "docx") in labels
    assert ("purchase_order", "standard", 1, "html") in labels
    assert ("purchase_order", "standard", 1, "docx") in labels


def test_adding_a_template_file_requires_no_code_change(tmp_path, monkeypatch):
    """Drop a file in the folder; it must show up and be renderable."""
    import shutil

    from app.config import get_settings

    settings = get_settings()
    root = tmp_path / "documents"
    shutil.copytree(settings.template_dir, root)
    (root / "invoice" / "compact-v3.html").write_text(
        "<html><body><h1>{{ document.title }} {{ document.number }}</h1>"
        "<p>{{ totals.grand_total.fmt }}</p></body></html>",
        encoding="utf-8",
    )
    monkeypatch.setattr(settings, "template_dir", root)

    names = templates.list_names("invoice")
    assert "compact" in names

    resolved = templates.resolve("invoice", "html", name="compact")
    assert resolved.version == 3
    html = render.render_html(resolved, GOLDEN_SAMPLE)
    assert "URJ/25-26/0042" in html
    assert "29,50,000.00" in html


def test_the_highest_version_wins_when_none_is_requested(tmp_path, monkeypatch):
    import shutil

    from app.config import get_settings

    settings = get_settings()
    root = tmp_path / "documents"
    shutil.copytree(settings.template_dir, root)
    (root / "invoice" / "standard-v2.html").write_text("<html>v2</html>", encoding="utf-8")
    monkeypatch.setattr(settings, "template_dir", root)

    assert templates.resolve("invoice", "html", name="standard").version == 2
    assert templates.resolve("invoice", "html", name="standard", version=1).version == 1


def test_a_missing_template_names_what_it_looked_for():
    with pytest.raises(templates.TemplateNotFound, match="nonexistent"):
        templates.resolve("invoice", "html", name="nonexistent")


@pytest.mark.parametrize("doc_type", ["invoice", "purchase_order"])
@pytest.mark.parametrize("kind", ["html", "docx"])
def test_every_shipped_template_validates_against_the_golden_sample(doc_type, kind):
    """This is the 'did I break it?' button, run in CI."""
    result = render.validate_template(doc_type, kind, "standard")
    assert result.ok, f"{doc_type}/{kind}: {result.undefined_names or result.error}"


def test_validation_reports_an_undefined_name_rather_than_failing_silently(
    tmp_path, monkeypatch
):
    import shutil

    from app.config import get_settings

    settings = get_settings()
    root = tmp_path / "documents"
    shutil.copytree(settings.template_dir, root)
    (root / "invoice" / "broken-v1.html").write_text(
        "<html>{{ totals.grand_total.fmt }} {{ typo_variable }}</html>", encoding="utf-8"
    )
    monkeypatch.setattr(settings, "template_dir", root)

    result = render.validate_template("invoice", "html", "broken")
    assert not result.ok
    assert "typo_variable" in result.undefined_names


def test_an_undefined_variable_renders_empty_rather_than_raising(tmp_path, monkeypatch):
    """A half-filled document beats a 500 for the person editing a template."""
    import shutil

    from app.config import get_settings

    settings = get_settings()
    root = tmp_path / "documents"
    shutil.copytree(settings.template_dir, root)
    (root / "invoice" / "loose-v1.html").write_text(
        "<html>[{{ not_a_real_variable }}][{{ also.not.real }}]</html>", encoding="utf-8"
    )
    monkeypatch.setattr(settings, "template_dir", root)

    html = render.render_html(templates.resolve("invoice", "html", name="loose"), GOLDEN_SAMPLE)
    assert "[]" in html


def test_both_formats_receive_the_same_context(
    db, company, tax_rates, gujarat_customer, order_factory
):
    """The .docx and the PDF cannot disagree, because they share one dict."""
    order = order_factory(gujarat_customer)
    invoice = invoice_service.create_tax_invoice(db, order, actor="t")
    invoice_service.issue(db, invoice, actor="t", formats=("pdf", "docx"))
    db.commit()

    from docx import Document

    text = "\n".join(p.text for p in Document(invoice.docx_path).paragraphs)
    for table in Document(invoice.docx_path).tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)

    assert invoice.number in text
    assert invoice.render_context_json["totals"]["grand_total"]["fmt"] in text


def test_money_reaches_templates_pre_formatted_so_they_need_no_arithmetic():
    line = GOLDEN_SAMPLE["lines"][0]
    assert line["unit_price"]["fmt"] == "12,50,000.00"
    assert line["unit_price"]["raw"] == "1250000"
    assert GOLDEN_SAMPLE["totals"]["grand_total"]["fmt"] == "29,50,000.00"


def test_a_pdf_is_actually_produced_and_is_a_pdf(db, company, tax_rates, gujarat_customer, order_factory):
    from pathlib import Path

    order = order_factory(gujarat_customer)
    invoice = invoice_service.create_tax_invoice(db, order, actor="t")
    invoice_service.issue(db, invoice, actor="t")
    db.commit()

    data = Path(invoice.pdf_path).read_bytes()
    assert data.startswith(b"%PDF")
    assert len(data) > 3000
